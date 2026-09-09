# -*- coding: utf-8 -*-
"""把 md + png 真题整理转成 docx（图片嵌入、中文文件名、LaTeX 保留）

用法：
    python build_docx.py <md目录> <docx输出目录>
    python build_docx.py "G:\\爱学\\图库对应测试题整理\\md和png" "G:\\爱学\\图库对应测试题整理"

依赖：python-docx
"""
import sys, os, re, argparse
from pathlib import Path

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

# 模板 ID → 中文名（可选；没配置就用 md 文件名）
CN_NAMES = {
    "quadratic-function":      "二次函数（顶点式）",
    "linear-function":         "一次函数 y=kx+b",
    "inverse-proportional":    "反比例函数（双曲线 y=k/x）",
    "exponential-function":    "指数函数 y=aˣ",
    "logarithm-function":     "对数函数 y=logₐx",
    "power-function-compare":  "幂函数图像对比",
    "trig-sine-wave":          "正弦型函数 y=A sin(ωx+φ)+k",
    "tangent-function":        "正切函数 y=tanx",
    "hook-function":           "对勾函数 y=x+a/x",
    "piecewise-function":      "分段与绝对值函数",
}

TEMPLATE_ORDER = list(CN_NAMES.keys())

# Markdown 图片引用 ![alt](path)
IMG_RE = re.compile(r'!\[([^\]]*)\]\(([^)]+)\)')
INLINE_LATEX_RE = re.compile(r'\$([^$\n]+?)\$')
DISPLAY_LATEX_RE = re.compile(r'\$\$(.+?)\$\$', re.DOTALL)


def set_font(run, chinese="SimSun", latin="Times New Roman", size=12, bold=False):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = latin
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        from docx.oxml import OxmlElement
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    rFonts.set(qn('w:eastAsia'), chinese)


def add_md_text(paragraph, text, base_size=11):
    """按 LaTeX 片段切分后写入 paragraph。行间公式另起居中段落。"""
    if not text.strip():
        return
    parts = DISPLAY_LATEX_RE.split(text)
    for i, seg in enumerate(parts):
        if not seg.strip():
            continue
        if i % 2 == 1:
            p2 = paragraph._parent.add_paragraph()
            p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p2.add_run(f"$${seg.strip()}$$")
            set_font(run, size=base_size)
            continue
        subparts = INLINE_LATEX_RE.split(seg)
        for j, sub in enumerate(subparts):
            if not sub:
                continue
            if j % 2 == 1:
                run = paragraph.add_run(f"${sub}$")
            else:
                run = paragraph.add_run(sub)
            set_font(run, size=base_size)


def parse_md(md_path: Path) -> dict:
    text = md_path.read_text(encoding="utf-8")
    lines = text.splitlines()
    tpl_id = md_path.stem
    cn_name = CN_NAMES.get(tpl_id, tpl_id)

    # 元信息（到第一个 --- 分隔线）
    meta_lines = []
    body_start = 0
    for i, ln in enumerate(lines):
        if ln.startswith("---"):
            body_start = i + 1
            break
        meta_lines.append(ln)

    desc = ""
    for ln in meta_lines:
        m = re.match(r"\*\*模板说明\*\*[：:]\s*(.+)", ln)
        if m:
            desc = m.group(1).strip()
            break
    if not desc and meta_lines:
        desc = " ".join(l.strip() for l in meta_lines[1:] if l.strip())[:200]

    rest = "\n".join(lines[body_start:])
    rest = re.sub(r'## 使用的试卷.*?(?=\n## )', '', rest, flags=re.DOTALL)

    q_blocks = re.split(r'(?=^## 真题)', rest, flags=re.MULTILINE)
    questions = []
    for blk in q_blocks:
        blk = blk.strip()
        if not blk or not blk.startswith("## 真题"):
            continue
        q = {"raw": blk}
        title_m = re.match(r'## 真题\s*[\d/\s]*[：:]?\s*(.+)', blk)
        if title_m:
            q["title"] = title_m.group(1).strip()
        for section in ["题型", "题干", "配图", "答案", "解析", "考点"]:
            pattern = rf'\*\*{section}\*\*[：:]\s*\n?(.*?)(?=\n\*\*|\n##|$)'
            m = re.search(pattern, blk, re.DOTALL)
            q[section] = m.group(1).strip() if m else ""
        img_m = IMG_RE.search(q.get("配图", ""))
        q["img_rel"] = img_m.group(2) if img_m else None
        questions.append(q)

    return {"tpl_id": tpl_id, "cn_name": cn_name, "desc": desc, "questions": questions}


def render_docx(info: dict, out_dir: Path, src_dir: Path) -> Path:
    doc = Document()
    style = doc.styles['Normal']
    style.font.name = 'Times New Roman'
    style.font.size = Pt(11)
    rPr = style.element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        from docx.oxml import OxmlElement
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    rFonts.set(qn('w:eastAsia'), 'SimSun')

    h = doc.add_heading(level=0)
    run = h.add_run(f"{info['cn_name']}")
    set_font(run, size=22, bold=True)
    h2 = doc.add_heading(level=1)
    run = h2.add_run("高考真题配图整理")
    set_font(run, size=14, bold=True)

    if info["desc"]:
        p = doc.add_paragraph()
        run = p.add_run("模板说明：")
        set_font(run, size=11, bold=True)
        run = p.add_run(info["desc"])
        set_font(run, size=11)
    doc.add_paragraph()

    for idx, q in enumerate(info["questions"], 1):
        doc.add_paragraph("─" * 50)

        hp = doc.add_heading(level=2)
        t = re.sub(r'\*\*', '', q.get("title", f"真题 {idx}")).strip()
        run = hp.add_run(f"{idx}. {t}")
        set_font(run, size=14, bold=True)

        if q.get("题型"):
            p = doc.add_paragraph()
            run = p.add_run("题型：")
            set_font(run, bold=True, size=11)
            run = p.add_run(re.sub(r'\*\*', '', q["题型"]))
            set_font(run, size=11)

        if q.get("题干"):
            p = doc.add_paragraph()
            run = p.add_run("题干：")
            set_font(run, bold=True, size=11)
            body = re.sub(r'^\*\*', '', q["题干"], flags=re.MULTILINE)
            body = re.sub(r'\*\*$', '', body.strip())
            for para_text in re.split(r'\n\s*\n', body):
                para_text = para_text.strip()
                if not para_text:
                    continue
                p2 = doc.add_paragraph()
                p2.paragraph_format.first_line_indent = Cm(0.74)
                add_md_text(p2, para_text, base_size=11)

        if q.get("img_rel"):
            img_path = None
            rel = q["img_rel"]
            candidates = [
                src_dir / rel,
                src_dir / info["tpl_id"] / Path(rel).name,
                src_dir / rel.replace("/", "\\"),
                src_dir / "md和png" / rel,
                src_dir / "md和png" / info["tpl_id"] / Path(rel).name,
            ]
            for c in candidates:
                if c.exists():
                    img_path = c
                    break
            p = doc.add_paragraph()
            run = p.add_run("配图：")
            set_font(run, bold=True, size=11)
            pimg = doc.add_paragraph()
            pimg.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if img_path:
                try:
                    pimg.add_run().add_picture(str(img_path), width=Cm(14))
                except Exception as ex:
                    pimg.add_run(f"[图片嵌入失败：{ex}]")
            else:
                pimg.add_run(f"[配图未找到: {rel}]")

        if q.get("答案"):
            p = doc.add_paragraph()
            run = p.add_run("答案：")
            set_font(run, bold=True, size=11)
            p2 = doc.add_paragraph()
            add_md_text(p2, re.sub(r'\*\*', '', q["答案"].strip()), base_size=11)

        if q.get("解析"):
            p = doc.add_paragraph()
            run = p.add_run("解析：")
            set_font(run, bold=True, size=11)
            body = re.sub(r'\*\*', '', q["解析"])
            body = re.sub(r'###\s*', '', body)
            for para_text in re.split(r'\n\s*\n', body):
                para_text = para_text.strip()
                if not para_text:
                    continue
                p2 = doc.add_paragraph()
                p2.paragraph_format.first_line_indent = Cm(0.74)
                add_md_text(p2, para_text, base_size=11)

        if q.get("考点"):
            p = doc.add_paragraph()
            run = p.add_run("考点：")
            set_font(run, bold=True, size=11)
            p2 = doc.add_paragraph()
            add_md_text(p2, re.sub(r'\*\*', '', q["考点"]), base_size=10.5)

    seq = TEMPLATE_ORDER.index(info["tpl_id"]) + 1 if info["tpl_id"] in TEMPLATE_ORDER else 99
    safe_cn = re.sub(r'[\\/:*?"<>|]', '_', info['cn_name'])
    out_name = f"{seq:02d}.{safe_cn}-真题整理.docx"
    out_path = out_dir / out_name
    out_dir.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    return out_path


def main():
    ap = argparse.ArgumentParser(description="md + png → docx")
    ap.add_argument("src", help="md 文件所在目录（png 子目录也在里面）")
    ap.add_argument("out", help="docx 输出目录")
    ap.add_argument("--glob", default="*.md", help="md 文件 glob，默认 *.md")
    args = ap.parse_args()

    src_dir = Path(args.src).resolve()
    out_dir = Path(args.out).resolve()
    if not src_dir.is_dir():
        print(f"[错误] md 目录不存在: {src_dir}"); sys.exit(1)

    md_files = sorted(src_dir.glob(args.glob))
    if not md_files:
        print(f"[错误] 没找到 {args.glob} 在 {src_dir}"); sys.exit(1)

    ok = fail = 0
    for md_path in md_files:
        try:
            info = parse_md(md_path)
            out = render_docx(info, out_dir, src_dir)
            print(f"[OK]  {out.name}  ({len(info['questions'])} 题)")
            ok += 1
        except Exception as ex:
            import traceback
            print(f"[FAIL] {md_path.name}: {ex}")
            traceback.print_exc()
            fail += 1
    print(f"\n完成：{ok} 成功，{fail} 失败 → {out_dir}")


if __name__ == "__main__":
    main()
